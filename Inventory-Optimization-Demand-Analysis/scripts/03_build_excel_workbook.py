"""
03_build_excel_workbook.py
==========================
Builds excel/inventory_analysis.xlsx - the Excel reporting layer.

Design notes
------------
1. The 73,100-row fact table stays in Power BI. Excel holds the aggregated
   reporting layer (100 store-SKUs, 25 months, 5 stores) plus the latest day's
   action register, which is what a workbook is actually good at.

2. Classification happens at POSITION-DAY level, not at SKU-average level.
   Averaging cover across 731 days washes the signal out - every store-SKU
   averages about 2.0 days of cover and looks healthy. The exposure is
   episodic, so the scorecard reports how OFTEN each position falls critical
   or excess, and the Action Register shows the latest day's actual status.

3. The scorecard carries a noise-band test. With 731 daily observations per
   position, the binomial standard error on an exception rate is about 1.45pp,
   so differences under roughly 4pp between positions are sampling noise.
   Testing that in the workbook keeps the analysis honest.

Run:  python scripts/03_build_excel_workbook.py
"""

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

OUT = "excel/inventory_analysis.xlsx"

sc = pd.read_csv("data/processed/sku_scorecard.csv").sort_values("SKU_Key").reset_index(drop=True)
monthly = pd.read_csv("data/processed/monthly_summary.csv")
fact = pd.read_csv("data/processed/fact_inventory_daily.csv", parse_dates=["Date"])
raw = pd.read_csv("data/raw/retail_store_inventory.csv", nrows=500)

latest_date = fact["Date"].max()
latest = fact[fact["Date"] == latest_date].sort_values("SKU_Key").reset_index(drop=True)

CRIT_RATE = 100 * (fact["Stock_Status"] == "Critical").mean()
EXC_RATE = 100 * (fact["Stock_Status"] == "Excess").mean()
N_OBS = fact["Date"].nunique()

FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
TITLE_FONT = Font(name=FONT, bold=True, size=14, color="1F3864")
SUB_FONT = Font(name=FONT, italic=True, size=9, color="595959")
BODY = Font(name=FONT, size=10)
BOLD = Font(name=FONT, size=10, bold=True)
INPUT_FONT = Font(name=FONT, size=10, bold=True, color="0000FF")
INPUT_FILL = PatternFill("solid", fgColor="FFFF00")
GREY = PatternFill("solid", fgColor="F2F2F2")
BAND_FILL = PatternFill("solid", fgColor="4472C4")
KPI_FONT = Font(name=FONT, bold=True, size=12)
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

MONEY = '$#,##0;($#,##0);-'
NUM2 = '#,##0.00'
INT = '#,##0'
PCT1 = '0.0"%"'
PCTF = '0.0%'


def style_header(ws, row, ncols, height=30):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill, cell.font = HDR_FILL, HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BOX
    ws.row_dimensions[row].height = height


def autosize(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


wb = Workbook()

# ===========================================================================
# SHEET 1 - README
# ===========================================================================
ws = wb.active
ws.title = "README"
ws["A1"] = "Inventory Optimization & Demand Analysis"
ws["A1"].font = Font(name=FONT, bold=True, size=16, color="1F3864")
ws["A2"] = "Excel reporting layer | Source: retail_store_inventory.csv | 73,100 rows | 2022-01-01 to 2024-01-01"
ws["A2"].font = SUB_FONT

lines = [
    ("", ""),
    ("SHEETS", ""),
    ("Parameters", "The only sheet you edit. Yellow cells with blue text are inputs. Changing a cover "
                   "threshold re-classifies the Action Register and the scorecard flags."),
    ("Action_Register", f"The decision list: all 100 stock positions as at {latest_date.date()}, each "
                        "with a live status and recommended action. This is the sheet a manager uses."),
    ("SKU_Scorecard", "One row per store-product. Reports how OFTEN each position falls critical or "
                      "excess across 731 days, with a noise-band test on whether the gap is real."),
    ("Monthly_Summary", "25 monthly periods. Complete Month = 0 marks a partial period."),
    ("Store_Summary", "Store roll-up via SUMIF / AVERAGEIF / COUNTIFS over the scorecard."),
    ("KPI_Dashboard", "Headline KPIs, all formula-driven."),
    ("Data_Dictionary", "Every source and derived field with its definition."),
    ("Raw_Sample", "First 500 source rows. Full file at data/raw/."),
    ("", ""),
    ("COLOUR LEGEND", ""),
    ("Blue text on yellow fill", "Input cell - safe to edit"),
    ("Black text", "Formula - do not overwrite"),
    ("Grey fill", "Aggregated fact imported from the Python build step"),
    ("", ""),
    ("WHY CLASSIFICATION IS PER DAY, NOT PER PRODUCT", ""),
    ("The exposure is episodic", "Averaged over 731 days every store-SKU sits near 2.0 days of cover and "
                                 "looks healthy. The risk appears on individual days, so the unit of "
                                 "decision is the position-day, not the product."),
    ("Noise band", f"Each position has {N_OBS} daily observations, giving a binomial standard error of "
                   "about 1.45pp on an exception rate. Gaps under roughly 4pp between positions are "
                   "sampling noise, and the scorecard says so rather than ranking them anyway."),
    ("", ""),
    ("KEY ASSUMPTIONS", ""),
    ("No cost column exists", "Inventory is valued at RETAIL price. Margin and cost-based turnover "
                              "cannot be computed from this dataset."),
    ("No reorder level column", "The reorder point is derived as Average Daily Demand x the Reorder "
                                "Cover parameter. It is a stated policy, not a source field."),
    ("No lead time or supplier", "Those KPIs are excluded rather than invented."),
    ("Category and Region", "Vary row-to-row within the same store-product, so they are transaction "
                            "labels, not product or store attributes. No decision rests on them."),
]
r = 4
for a, b in lines:
    ws.cell(row=r, column=1, value=a).font = Font(name=FONT, bold=bool(a and not b), size=10)
    c = ws.cell(row=r, column=2, value=b)
    c.font, c.alignment = BODY, Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 30 if len(b) > 90 else 15
    r += 1
autosize(ws, [30, 98])

# ===========================================================================
# SHEET 2 - PARAMETERS
# ===========================================================================
ws = wb.create_sheet("Parameters")
ws["A1"] = "Policy Parameters"
ws["A1"].font = TITLE_FONT
ws["A2"] = "Edit the yellow cells. Every classification in this workbook reads from here."
ws["A2"].font = SUB_FONT

params = [
    ("Critical cover threshold (days)", 1.0,
     "Below this, stock covers under one day of demand. Action = REPLENISH NOW."),
    ("Reorder cover threshold (days)", 1.5,
     "Reorder point in days of demand. Set near the 25th percentile of observed cover (1.19 days)."),
    ("Target cover (days)", 3.0,
     "Stock above this is excess. Set near the 75th percentile of observed cover (2.84 days)."),
    ("Analysis period (days)", int(N_OBS),
     "Days in the source panel: 2022-01-01 to 2024-01-01 inclusive."),
    ("Days per year", 365.25, "Used to annualise inventory turnover."),
    ("System critical rate (%)", round(CRIT_RATE, 2),
     "Share of all 73,100 position-days classified Critical. Centre of the noise band."),
    ("System excess rate (%)", round(EXC_RATE, 2),
     "Share of all 73,100 position-days classified Excess. Centre of the noise band."),
    ("Noise band width (std errors)", 2.8,
     "Largest deviation expected from chance across 100 positions. Beyond this, a gap is real."),
]
ws["A4"], ws["B4"], ws["C4"] = "Parameter", "Value", "Basis"
style_header(ws, 4, 3)
for i, (name, val, note) in enumerate(params):
    row = 5 + i
    ws.cell(row=row, column=1, value=name).font = BODY
    c = ws.cell(row=row, column=2, value=val)
    c.font, c.fill, c.border, c.number_format = INPUT_FONT, INPUT_FILL, BOX, NUM2
    n = ws.cell(row=row, column=3, value=note)
    n.font, n.alignment = BODY, Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[row].height = 28

ws["A14"] = "Derived noise band (calculated)"
ws["A14"].font = BOLD
for i, (lbl, f) in enumerate([
    ("Critical std error per position (pp)", "=SQRT(($B$10/100)*(1-$B$10/100)/$B$8)*100"),
    ("Critical rate lower bound (%)", "=$B$10-$B$12*$B$15"),
    ("Critical rate upper bound (%)", "=$B$10+$B$12*$B$15"),
    ("Excess std error per position (pp)", "=SQRT(($B$11/100)*(1-$B$11/100)/$B$8)*100"),
    ("Excess rate lower bound (%)", "=$B$11-$B$12*$B$18"),
    ("Excess rate upper bound (%)", "=$B$11+$B$12*$B$18"),
]):
    row = 15 + i
    ws.cell(row=row, column=1, value=lbl).font = BODY
    c = ws.cell(row=row, column=2, value=f)
    c.font, c.border, c.number_format = BODY, BOX, NUM2
autosize(ws, [36, 13, 80])

P_CRIT, P_REORD, P_TARGET, P_DAYS, P_YR = "$B$5", "$B$6", "$B$7", "$B$8", "$B$9"
P_CLO, P_CHI, P_ELO, P_EHI = "$B$16", "$B$17", "$B$19", "$B$20"

# ===========================================================================
# SHEET 3 - ACTION REGISTER
# ===========================================================================
ws = wb.create_sheet("Action_Register")
ws["A1"] = f"Action Register - Stock Position as at {latest_date.date()}"
ws["A1"].font = TITLE_FONT
ws["A2"] = ("All 100 store-product positions on the final day of the panel. Status and action are live "
            "formulas: change a threshold on Parameters and this list re-prioritises.")
ws["A2"].font = SUB_FONT

ah = ["SKU Key", "Store ID", "Product ID", "Inventory (units)", "Avg Daily Demand",
      "Days of Cover", "Reorder Point (units)", "Order Qty to Target (units)",
      "Units Actually Ordered", "Order Gap (units)", "Inventory Value ($)",
      "Stock Status", "Recommended Action", "Unit Price ($)"]
HR = 4
for i, h in enumerate(ah, start=1):
    ws.cell(row=HR, column=i, value=h)
style_header(ws, HR, len(ah), height=42)

for i, rec in latest.iterrows():
    r = HR + 1 + i
    for cidx, v in zip([1, 2, 3, 4, 5, 9, 14],
                       [rec["SKU_Key"], rec["Store_ID"], rec["Product_ID"],
                        int(rec["Inventory_Level"]), round(rec["Avg_Daily_Demand"], 2),
                        int(rec["Units_Ordered"]), round(rec["Price"], 2)]):
        c = ws.cell(row=r, column=cidx, value=v)
        c.font, c.fill, c.border = BODY, GREY, BOX
        if cidx in (4, 9):
            c.number_format = INT
        elif cidx in (5, 14):
            c.number_format = NUM2
    ws.cell(row=r, column=6, value=f"=IFERROR(D{r}/E{r},0)")
    ws.cell(row=r, column=7, value=f"=E{r}*Parameters!{P_REORD}")
    ws.cell(row=r, column=8, value=f"=MAX(0,E{r}*Parameters!{P_TARGET}-D{r})")
    ws.cell(row=r, column=10, value=f"=I{r}-H{r}")
    ws.cell(row=r, column=11, value=f"=D{r}*N{r}")
    ws.cell(row=r, column=12,
            value=(f'=IF(F{r}<Parameters!{P_CRIT},"Critical",'
                   f'IF(F{r}<Parameters!{P_REORD},"Below Reorder",'
                   f'IF(F{r}<=Parameters!{P_TARGET},"Healthy","Excess")))'))
    ws.cell(row=r, column=13,
            value=(f'=IF(L{r}="Critical","REPLENISH NOW",'
                   f'IF(L{r}="Below Reorder","REPLENISH",'
                   f'IF(L{r}="Healthy","MAINTAIN","REDUCE STOCK")))'))
    for cidx in [6, 7, 8, 10, 11, 12, 13]:
        c = ws.cell(row=r, column=cidx)
        c.font, c.border = BODY, BOX
        if cidx in (6, 7, 8, 10):
            c.number_format = NUM2
        elif cidx == 11:
            c.number_format = MONEY

ALAST = HR + len(latest)
tbl = Table(displayName="ActionRegister", ref=f"A{HR}:N{ALAST}")
tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
ws.add_table(tbl)
ws.freeze_panes = f"D{HR + 1}"
autosize(ws, [12, 10, 11, 15, 15, 13, 16, 17, 16, 14, 16, 14, 18, 13])

# ===========================================================================
# SHEET 4 - SKU SCORECARD
# ===========================================================================
ws = wb.create_sheet("SKU_Scorecard")
ws["A1"] = "Store-SKU Scorecard - Exception Frequency Across 731 Days"
ws["A1"].font = TITLE_FONT
ws["A2"] = ("Grey columns are aggregated facts; white columns are live formulas. The 'Beyond Noise?' "
            "columns test whether a position's exception rate genuinely differs from the system rate "
            "or is just sampling variation.")
ws["A2"].font = SUB_FONT

headers = ["SKU Key", "Store ID", "Product ID", "Total Units Sold", "Avg Daily Demand",
           "Demand Std Dev", "Avg Inventory (units)", "Avg Price ($)", "Days Critical",
           "Days Excess", "Sold-Out Days", "Stockout-Risk Days", "Unserved Units",
           "Forecast Bias (units)", "Inventory Turnover (x/yr)", "Avg Days of Cover",
           "Avg Inventory Value ($)", "% Days Critical", "% Days Excess",
           "Critical Beyond Noise?", "Excess Beyond Noise?", "Monitoring Priority"]
HDR_ROW = 4
for i, h in enumerate(headers, start=1):
    ws.cell(row=HDR_ROW, column=i, value=h)
style_header(ws, HDR_ROW, len(headers), height=44)

sc["Avg_Price"] = (sc["Avg_Inventory_Value"] / sc["Avg_Inventory"]).round(2)

for i, rec in sc.iterrows():
    r = HDR_ROW + 1 + i
    facts = [rec["SKU_Key"], rec["Store_ID"], rec["Product_ID"], int(rec["Total_Units_Sold"]),
             rec["Avg_Daily_Demand"], rec["Demand_StdDev"], rec["Avg_Inventory"], rec["Avg_Price"],
             int(rec["Days_Critical"]), int(rec["Days_Excess"]), int(rec["Sold_Out_Days"]),
             int(rec["Stockout_Risk_Days"]), round(rec["Unserved_Units"], 1),
             round(rec["Forecast_Bias"], 2)]
    for cidx, v in enumerate(facts, start=1):
        c = ws.cell(row=r, column=cidx, value=v)
        c.font, c.fill, c.border = BODY, GREY, BOX
        if cidx in (4, 9, 10, 11, 12):
            c.number_format = INT
        elif cidx in (5, 6, 7, 8, 13, 14):
            c.number_format = NUM2

    ws.cell(row=r, column=15,
            value=f"=IFERROR(D{r}/G{r}/(Parameters!{P_DAYS}/Parameters!{P_YR}),0)")
    ws.cell(row=r, column=16, value=f"=IFERROR(G{r}/E{r},0)")
    ws.cell(row=r, column=17, value=f"=G{r}*H{r}")
    ws.cell(row=r, column=18, value=f"=100*I{r}/Parameters!{P_DAYS}")
    ws.cell(row=r, column=19, value=f"=100*J{r}/Parameters!{P_DAYS}")
    ws.cell(row=r, column=20,
            value=f'=IF(OR(R{r}<Parameters!{P_CLO},R{r}>Parameters!{P_CHI}),"YES","no - within noise")')
    ws.cell(row=r, column=21,
            value=f'=IF(OR(S{r}<Parameters!{P_ELO},S{r}>Parameters!{P_EHI}),"YES","no - within noise")')
    ws.cell(row=r, column=22,
            value=(f'=IF(AND(T{r}="YES",R{r}>Parameters!{P_CHI}),"Watch - stockout",'
                   f'IF(AND(U{r}="YES",S{r}>Parameters!{P_EHI}),"Watch - excess capital","Routine"))'))
    for cidx in range(15, 23):
        c = ws.cell(row=r, column=cidx)
        c.font, c.border = BODY, BOX
        if cidx in (15, 16):
            c.number_format = NUM2
        elif cidx == 17:
            c.number_format = MONEY
        elif cidx in (18, 19):
            c.number_format = PCT1

LAST = HDR_ROW + len(sc)
tbl = Table(displayName="SKUScorecard", ref=f"A{HDR_ROW}:V{LAST}")
tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
ws.add_table(tbl)
ws.freeze_panes = f"D{HDR_ROW + 1}"
autosize(ws, [12, 10, 11, 15, 14, 13, 15, 12, 12, 12, 12, 14, 13, 14,
              16, 14, 17, 13, 13, 17, 17, 20])

# ===========================================================================
# SHEET 5 - MONTHLY SUMMARY
# ===========================================================================
ws = wb.create_sheet("Monthly_Summary")
ws["A1"] = "Monthly Demand and Inventory Summary"
ws["A1"].font = TITLE_FONT
ws["A2"] = ("Complete Month = 0 marks a partial period. January 2024 holds a single day (2024-01-01) "
            "and is excluded from growth calculations.")
ws["A2"].font = SUB_FONT

mh = ["Year-Month", "Units Sold", "Revenue ($)", "Avg Inventory (units)", "Avg Days of Cover",
      "Critical Position-Days", "Excess Position-Days", "Avg Excess Value ($)",
      "Forecast Bias (units)", "Position Days", "Complete Month", "MoM Units Growth"]
for i, h in enumerate(mh, start=1):
    ws.cell(row=4, column=i, value=h)
style_header(ws, 4, len(mh))

for i, rec in monthly.iterrows():
    r = 5 + i
    vals = [rec["Year_Month"], int(rec["Units_Sold"]), rec["Revenue"], rec["Avg_Inventory"],
            rec["Avg_Days_of_Cover"], int(rec["Critical_Positions"]), int(rec["Excess_Positions"]),
            rec["Excess_Value"], rec["Forecast_Bias"], int(rec["Position_Days"]),
            int(rec["Is_Complete_Month"])]
    for cidx, v in enumerate(vals, start=1):
        c = ws.cell(row=r, column=cidx, value=v)
        c.font, c.border = BODY, BOX
        if cidx in (2, 6, 7, 10):
            c.number_format = INT
        elif cidx in (3, 8):
            c.number_format = MONEY
        elif cidx in (4, 5, 9):
            c.number_format = NUM2
    if i == 0:
        ws.cell(row=r, column=12, value="")
    else:
        ws.cell(row=r, column=12,
                value=f'=IF(OR(K{r}=0,K{r-1}=0),"",IFERROR(B{r}/B{r-1}-1,""))')
    gc = ws.cell(row=r, column=12)
    gc.font, gc.border, gc.number_format = BODY, BOX, PCTF

MLAST = 4 + len(monthly)
ws.freeze_panes = "A5"
autosize(ws, [12, 12, 15, 18, 16, 18, 17, 17, 16, 13, 14, 16])

# ===========================================================================
# SHEET 6 - STORE SUMMARY
# ===========================================================================
ws = wb.create_sheet("Store_Summary")
ws["A1"] = "Store Roll-Up"
ws["A1"].font = TITLE_FONT
ws["A2"] = "Every figure is a SUMIF / AVERAGEIF / COUNTIFS against SKU_Scorecard."
ws["A2"].font = SUB_FONT

sh = ["Store ID", "SKU Positions", "Total Units Sold", "Avg Inventory Value ($)",
      "Avg Days of Cover", "Avg Turnover (x/yr)", "Avg % Days Critical", "Avg % Days Excess",
      "Total Unserved Units", "Positions Flagged for Watch"]
for i, h in enumerate(sh, start=1):
    ws.cell(row=4, column=i, value=h)
style_header(ws, 4, len(sh), height=40)

SR = f"SKU_Scorecard!$B${HDR_ROW + 1}:$B${LAST}"
for i, store in enumerate(sorted(sc["Store_ID"].unique())):
    r = 5 + i
    ws.cell(row=r, column=1, value=store).font = BODY
    ws.cell(row=r, column=2, value=f'=COUNTIF({SR},$A{r})')
    ws.cell(row=r, column=3, value=f'=SUMIF({SR},$A{r},SKU_Scorecard!$D${HDR_ROW+1}:$D${LAST})')
    ws.cell(row=r, column=4, value=f'=SUMIF({SR},$A{r},SKU_Scorecard!$Q${HDR_ROW+1}:$Q${LAST})')
    ws.cell(row=r, column=5, value=f'=AVERAGEIF({SR},$A{r},SKU_Scorecard!$P${HDR_ROW+1}:$P${LAST})')
    ws.cell(row=r, column=6, value=f'=AVERAGEIF({SR},$A{r},SKU_Scorecard!$O${HDR_ROW+1}:$O${LAST})')
    ws.cell(row=r, column=7, value=f'=AVERAGEIF({SR},$A{r},SKU_Scorecard!$R${HDR_ROW+1}:$R${LAST})')
    ws.cell(row=r, column=8, value=f'=AVERAGEIF({SR},$A{r},SKU_Scorecard!$S${HDR_ROW+1}:$S${LAST})')
    ws.cell(row=r, column=9, value=f'=SUMIF({SR},$A{r},SKU_Scorecard!$M${HDR_ROW+1}:$M${LAST})')
    ws.cell(row=r, column=10,
            value=f'=COUNTIFS({SR},$A{r},SKU_Scorecard!$V${HDR_ROW+1}:$V${LAST},"Watch*")')
    for cidx in range(2, 11):
        c = ws.cell(row=r, column=cidx)
        c.font, c.border = BODY, BOX
        if cidx == 4:
            c.number_format = MONEY
        elif cidx in (5, 6, 9):
            c.number_format = NUM2
        elif cidx in (7, 8):
            c.number_format = PCT1
        else:
            c.number_format = INT
autosize(ws, [11, 15, 17, 21, 17, 18, 18, 18, 19, 24])

# ===========================================================================
# SHEET 7 - KPI DASHBOARD
# ===========================================================================
ws = wb.create_sheet("KPI_Dashboard")
ws["A1"] = "Headline KPIs"
ws["A1"].font = TITLE_FONT
ws["A2"] = "All values are formulas over the other sheets. Nothing is hardcoded."
ws["A2"].font = SUB_FONT

AR_STATUS = f"Action_Register!$L${HR+1}:$L${ALAST}"
SCP = "SKU_Scorecard!"

kpis = [
    ("PERIOD TOTALS", None, None, None),
    ("Total Units Sold", f"=SUM(Monthly_Summary!$B$5:$B${MLAST})", INT,
     "All 100 positions across 731 days."),
    ("Total Revenue", f"=SUM(Monthly_Summary!$C$5:$C${MLAST})", MONEY,
     "Units sold x price x (1 - discount)."),
    ("Total Inventory Value on Hand", f"=SUM({SCP}$Q${HDR_ROW+1}:$Q${LAST})", MONEY,
     "Average daily value across all positions, at retail. No cost column exists."),
    ("INVENTORY EFFICIENCY", None, None, None),
    ("Avg Days of Cover", f"=AVERAGE({SCP}$P${HDR_ROW+1}:$P${LAST})", NUM2,
     "Average inventory divided by average daily demand."),
    ("Avg Inventory Turnover (x/yr)", f"=AVERAGE({SCP}$O${HDR_ROW+1}:$O${LAST})", NUM2,
     "Annualised. High because cover sits near two days - a fast-moving retail profile."),
    ("System Critical Rate (%)", f"=AVERAGE({SCP}$R${HDR_ROW+1}:$R${LAST})", PCT1,
     "Share of position-days below one day of cover."),
    ("System Excess Rate (%)", f"=AVERAGE({SCP}$S${HDR_ROW+1}:$S${LAST})", PCT1,
     "Share of position-days above target cover."),
    ("SERVICE RISK", None, None, None),
    ("Total Sold-Out Days", f"=SUM({SCP}$K${HDR_ROW+1}:$K${LAST})", INT,
     "Position-days that ended with zero stock; true demand was censored."),
    ("Total Stockout-Risk Days", f"=SUM({SCP}$L${HDR_ROW+1}:$L${LAST})", INT,
     "Position-days where forecast demand exceeded stock on hand."),
    ("Total Unserved Units", f"=SUM({SCP}$M${HDR_ROW+1}:$M${LAST})", INT,
     "Forecast demand the position could not have covered."),
    ("FORECAST QUALITY", None, None, None),
    ("Avg Forecast Bias (units/day)", f"=AVERAGE({SCP}$N${HDR_ROW+1}:$N${LAST})", NUM2,
     "Positive means the supplied forecast runs above what actually sold."),
    ("LATEST-DAY ACTION LIST", None, None, None),
    ("Positions: REPLENISH NOW", f'=COUNTIF({AR_STATUS},"Critical")', INT,
     f"Below one day of cover on {latest_date.date()}."),
    ("Positions: REPLENISH", f'=COUNTIF({AR_STATUS},"Below Reorder")', INT,
     "Between the critical and reorder thresholds."),
    ("Positions: MAINTAIN", f'=COUNTIF({AR_STATUS},"Healthy")', INT,
     "Between reorder and target cover."),
    ("Positions: REDUCE STOCK", f'=COUNTIF({AR_STATUS},"Excess")', INT, "Above target cover."),
    ("Units to Order Today", f"=SUM(Action_Register!$H${HR+1}:$H${ALAST})", INT,
     "Units required to bring every position up to target cover."),
    ("Today's Order Gap (units)", f"=SUM(Action_Register!$J${HR+1}:$J${ALAST})", INT,
     "Actual order quantity minus requirement. Negative means under-ordering versus need."),
    ("MONITORING", None, None, None),
    ("Positions Flagged for Watch", f'=COUNTIF({SCP}$V${HDR_ROW+1}:$V${LAST},"Watch*")', INT,
     "Positions whose exception rate is genuinely outside the noise band."),
    ("Positions Classed Routine", f'=COUNTIF({SCP}$V${HDR_ROW+1}:$V${LAST},"Routine")', INT,
     "Exception rate indistinguishable from the system average."),
]

ws["A4"], ws["B4"], ws["C4"] = "KPI", "Value", "Definition"
style_header(ws, 4, 3)
r = 5
for name, formula, fmt, note in kpis:
    if formula is None:
        for cc in range(1, 4):
            cell = ws.cell(row=r, column=cc, value=name if cc == 1 else None)
            cell.fill, cell.border = BAND_FILL, BOX
            cell.font = Font(name=FONT, bold=True, size=10, color="FFFFFF")
        r += 1
        continue
    ws.cell(row=r, column=1, value=name).font = BOLD
    c = ws.cell(row=r, column=2, value=formula)
    c.font, c.number_format, c.border = KPI_FONT, fmt, BOX
    c.alignment = Alignment(horizontal="right")
    n = ws.cell(row=r, column=3, value=note)
    n.font, n.alignment = BODY, Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 26
    r += 1
autosize(ws, [32, 18, 78])

# ===========================================================================
# SHEET 8 - DATA DICTIONARY
# ===========================================================================
ws = wb.create_sheet("Data_Dictionary")
ws["A1"] = "Data Dictionary"
ws["A1"].font = TITLE_FONT
for i, h in enumerate(["Field", "Origin", "Type", "Definition / Note"], start=1):
    ws.cell(row=3, column=i, value=h)
style_header(ws, 3, 4)

DICT = [
    ("Date", "Source", "Date", "2022-01-01 to 2024-01-01. 731 consecutive days, no gaps."),
    ("Store ID", "Source", "Text", "S001-S005. Five stores."),
    ("Product ID", "Source", "Text", "P0001-P0020. Twenty products, each stocked in all five stores."),
    ("Category", "Source", "Text",
     "Five values. WARNING: varies row-to-row for the same product, so it is a transaction label, not a product attribute."),
    ("Region", "Source", "Text",
     "Four values. WARNING: varies row-to-row for the same store, so it is not a store attribute."),
    ("Inventory Level", "Source", "Integer", "Units on hand, daily snapshot."),
    ("Units Sold", "Source", "Integer", "Units sold that day. Never exceeds Inventory Level."),
    ("Units Ordered", "Source", "Integer",
     "Units ordered that day, 20-200. Uncorrelated with demand or stock position (|r| < 0.002)."),
    ("Demand Forecast", "Source", "Decimal",
     "Supplied forecast. Carries a systematic upward bias of about 5 units."),
    ("Price", "Source", "Decimal", "Shelf price per unit, roughly $10-$100."),
    ("Discount", "Source", "Integer", "Whole-number percent: 0, 5, 10, 15 or 20."),
    ("Weather Condition", "Source", "Text", "Four values. No measurable effect on units sold."),
    ("Holiday/Promotion", "Source", "Binary",
     "0 or 1. No statistically significant volume effect (p = 0.92)."),
    ("Competitor Pricing", "Source", "Decimal", "Tracks Price almost exactly (r = 0.99)."),
    ("Seasonality", "Source", "Text", "Four values, not aligned to the calendar month."),
    ("SKU_Key", "Derived", "Text", "Store ID + Product ID. The true stock position key."),
    ("Net_Price", "Derived", "Decimal", "Price x (1 - Discount/100)."),
    ("Revenue", "Derived", "Decimal", "Units Sold x Net Price."),
    ("Inventory_Value", "Derived", "Decimal",
     "Inventory Level x Price. At retail - no cost column exists."),
    ("Avg_Daily_Demand", "Derived", "Decimal", "Trailing 28-day mean of Units Sold per SKU."),
    ("Days_of_Cover", "Derived", "Decimal", "Inventory Level / Avg Daily Demand. The core metric."),
    ("Reorder_Point", "Derived", "Decimal",
     "Avg Daily Demand x reorder cover parameter. A stated policy, not a source field."),
    ("Excess_Units", "Derived", "Decimal", "Units held above target cover, floored at zero."),
    ("Excess_Value", "Derived", "Decimal", "Excess Units x Price."),
    ("Sold_Out", "Derived", "Binary", "1 when Units Sold equals Inventory Level - demand was censored."),
    ("Stockout_Risk", "Derived", "Binary", "1 when Demand Forecast exceeds Inventory Level."),
    ("Unserved_Units", "Derived", "Decimal", "Demand Forecast minus Inventory Level, floored at zero."),
    ("Forecast_Error", "Derived", "Decimal", "Demand Forecast minus Units Sold."),
    ("Order_Requirement", "Derived", "Decimal", "Units needed to reach target cover, floored at zero."),
    ("Order_Gap", "Derived", "Decimal", "Units Ordered minus Order Requirement."),
    ("Stock_Status", "Derived", "Text",
     "Critical / Below Reorder / Healthy / Excess, from days of cover."),
    ("Recommended_Action", "Derived", "Text",
     "REPLENISH NOW / REPLENISH / MAINTAIN / REDUCE STOCK."),
]
for i, (f, o, t, d) in enumerate(DICT):
    r = 4 + i
    for cidx, v in enumerate([f, o, t, d], start=1):
        c = ws.cell(row=r, column=cidx, value=v)
        c.font, c.border = BODY, BOX
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 26 if len(d) > 70 else 15
autosize(ws, [21, 11, 10, 84])

# ===========================================================================
# SHEET 9 - RAW SAMPLE
# ===========================================================================
ws = wb.create_sheet("Raw_Sample")
ws["A1"] = "Source Data - First 500 Rows"
ws["A1"].font = TITLE_FONT
ws["A2"] = "Reference only. The full 73,100-row file is at data/raw/retail_store_inventory.csv."
ws["A2"].font = SUB_FONT
for i, h in enumerate(raw.columns, start=1):
    ws.cell(row=4, column=i, value=h)
style_header(ws, 4, len(raw.columns))
for i, rec in raw.iterrows():
    for cidx, v in enumerate(rec.values, start=1):
        ws.cell(row=5 + i, column=cidx, value=v).font = BODY
ws.freeze_panes = "A5"
autosize(ws, [12, 10, 11, 13, 10, 15, 11, 13, 15, 9, 10, 17, 16, 17, 12])

wb.save(OUT)
print(f"Wrote {OUT}")
print("Sheets:", wb.sheetnames)
